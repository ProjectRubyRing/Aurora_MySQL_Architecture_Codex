"""Validate the deliverables against captured evidence, without connecting to a DB."""
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile

import openpyxl

from report_sources import SOURCE_MAP

ROOT = Path(__file__).resolve().parents[1]
VALIDATION = ROOT / 'validation'
VALIDATION.mkdir(exist_ok=True)
STEM = 'Aurora_GlobalDatabase_Design_Assessment_2026-10-11'
XLSX = ROOT / 'deliverables' / (STEM + '.xlsx')
MD = ROOT / 'deliverables' / (STEM + '.md')
RESULT = VALIDATION / 'validation_report.json'
checks = []


def check(name, condition, detail):
    checks.append({'check': name, 'result': 'PASS' if condition else 'FAIL', 'detail': detail})


def read_json(path):
    return json.loads(path.read_text(encoding='utf-8'))


def gray(color):
    if color is None or color.type != 'rgb':
        return True
    value = color.rgb[-6:]
    return value[0:2] == value[2:4] == value[4:6]


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=False)
    cached = openpyxl.load_workbook(XLSX, data_only=True)
    sections = read_json(ROOT / 'evidence' / 'report_sections.json')
    expected_sheets = {s['sheet'] for s in sections} | {'00_はじめに', '02_構成図', '15_接続数試算'}
    check('Workbook sheet completeness', set(wb.sheetnames) == expected_sheets,
          {'actual_count': len(wb.sheetnames), 'names': wb.sheetnames})
    differences = []
    style_errors = []
    refs = set()
    for section in sections:
        ws = wb[section['sheet']]
        for r, row in enumerate(section['rows'], 6):
            for c, expected in enumerate(row, 1):
                actual = ws.cell(r, c)
                if actual.value != str(expected):
                    differences.append(f'{ws.title}!{actual.coordinate}')
                if not actual.alignment.wrap_text:
                    style_errors.append(f'{ws.title}!{actual.coordinate}: not wrapped')
                refs.update(re.findall(r'\bS\d{2}\b', str(expected)))
        check(f'Table navigation: {ws.title}',
              bool(ws.freeze_panes and ws.auto_filter.ref and ws.print_area and ws.print_title_rows),
              {'rows': len(section['rows']), 'filter': ws.auto_filter.ref, 'freeze': ws.freeze_panes})
    check('Every table cell matches the common content source', not differences, differences)
    check('Wrapped body cells', not style_errors, style_errors)
    check('All citation IDs resolve', refs.issubset(SOURCE_MAP), sorted(refs - set(SOURCE_MAP)))

    for scope, count, defaults_sheet, recs_sheet in [
        ('cluster', 430, '20_クラスタ既定', '21_クラスタ推奨'),
        ('instance', 299, '22_インスタンス既定', '23_インスタンス推奨'),
    ]:
        raw = read_json(ROOT / 'evidence' / f'ap-northeast-1_{scope}_group.json')['Parameters']
        osaka = read_json(ROOT / 'evidence' / f'ap-northeast-3_{scope}_group.json')['Parameters']
        names = [p['ParameterName'] for p in raw]
        actual = [wb[defaults_sheet].cell(i, 2).value for i in range(6, 6 + count)]
        actual_recs = [wb[recs_sheet].cell(i, 2).value for i in range(6, 6 + count)]
        check(f'{scope}: full parameter set and order', len(raw) == count and names == actual == actual_recs,
              {'api_rows': len(raw), 'distinct_names': len(set(names))})
        mismatch = []
        omitted = 0
        for i, parameter in enumerate(raw, 6):
            value = parameter.get('ParameterValue', 'engine-default（APIは具体値を省略）')
            if value == '':
                value = '空文字（API返値）'
            if wb[defaults_sheet].cell(i, 4).value != value:
                mismatch.append(parameter['ParameterName'])
            omitted += 'ParameterValue' not in parameter
        check(f'{scope}: API default fidelity; omissions not invented', not mismatch,
              {'omitted_defaults': omitted, 'mismatch': mismatch})
        excluded = [p['ParameterName'] for p in raw if p.get('MinimumEngineVersion')]
        check(f'{scope}: later-version parameter explicitly excluded', excluded == ['aurora_transaction_timeout']
              and all('8.4.7対象外' in wb[defaults_sheet].cell(i, 6).value
                      for i, p in enumerate(raw, 6) if p.get('MinimumEngineVersion')),
              excluded)
        fields = ['ParameterName', 'ParameterValue', 'IsModifiable', 'AllowedValues', 'ApplyType', 'DataType', 'MinimumEngineVersion']
        left = [{k: p.get(k) for k in fields} for p in raw]
        right = [{k: p.get(k) for k in fields} for p in osaka]
        check(f'{scope}: Tokyo/Osaka default and constraint agreement', left == right,
              'SupportedEngineModes omissions are evaluated separately from service availability.')

    long_values = {}
    for row in wb['25_長い許容値'].iter_rows(min_row=6, values_only=True):
        long_values.setdefault(row[0], []).extend(row[2].split('\n'))
    raw_cluster = read_json(ROOT / 'evidence' / 'ap-northeast-1_cluster_group.json')['Parameters']
    reconstructed = {name: ','.join(values) for name, values in long_values.items()}
    expected_long = {p['ParameterName']: p['AllowedValues'] for p in raw_cluster if len(p.get('AllowedValues', '')) > 2000}
    check('Long allowed-value lists have no omissions or reordering', reconstructed == expected_long,
          {name: len(values) for name, values in long_values.items()})
    for region in ['ap-northeast-1', 'ap-northeast-3']:
        engine = read_json(ROOT / 'evidence' / f'{region}_engine.json')['DBEngineVersions'][0]
        orderable = read_json(ROOT / 'evidence' / f'{region}_serverless.json')
        check(f'{region}: 8.4.7 / db.serverless / Global availability',
              engine['EngineVersion'] == '8.4.mysql_aurora.8.4.7' and engine['Status'] == 'available'
              and engine['SupportsGlobalDatabases'] and any(x['DBInstanceClass'] == 'db.serverless'
              and x['SupportsGlobalDatabases'] for x in orderable),
              {'engine': engine['EngineVersion'], 'serverless_options': len(orderable)})
    native = next(p for p in raw_cluster if p['ParameterName'] == 'mysql_native_password')
    check('Aurora native password evidence', native['ParameterValue'] == 'ON' and not native['IsModifiable'],
          'Aurora-specific fact, distinct from Community MySQL 8.4.')

    font_errors, color_errors, height_errors, formula_errors, link_errors = [], [], [], [], []
    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.value is None:
                    continue
                if cell.font.name != 'Meiryo UI':
                    font_errors.append(f'{ws.title}!{cell.coordinate}:{cell.font.name}')
                if not gray(cell.font.color) or not gray(cell.fill.fgColor):
                    color_errors.append(f'{ws.title}!{cell.coordinate}')
                if cell.data_type == 'f' and ws.title != '15_接続数試算':
                    formula_errors.append(f'{ws.title}!{cell.coordinate}')
                if cell.hyperlink and cell.hyperlink.location and not cell.hyperlink.target:
                    dest = cell.hyperlink.location.split('!')[0].strip("'")
                    if dest not in wb.sheetnames:
                        link_errors.append(cell.hyperlink.location)
        for row_number, dim in ws.row_dimensions.items():
            if dim.height and dim.height > 409:
                height_errors.append(f'{ws.title}!{row_number}:{dim.height}')
    check('Meiryo UI on all populated cells', not font_errors, font_errors)
    check('Monochrome cell font and fill colors', not color_errors, color_errors)
    check('Excel row-height limit respected', not height_errors, height_errors)
    check('No unintended spreadsheet formulas', not formula_errors, formula_errors)
    check('Workbook internal links resolve', not link_errors, link_errors)
    expected_calculations = [1208, 1571, 729, 2400, 550, '枠内（性能は未検証）']
    actual_calculations = [cached['15_接続数試算'].cell(r, 3).value for r in range(19, 25)]
    check('Capacity formula cached results', actual_calculations == expected_calculations, actual_calculations)
    with zipfile.ZipFile(XLSX) as archive:
        archive_error = archive.testzip()
        xml_files = [name for name in archive.namelist() if name.endswith('.xml')]
        for name in xml_files:
            ET.fromstring(archive.read(name))
        check('XLSX ZIP CRC and XML parsing', archive_error is None, {'xml_parts': len(xml_files), 'crc_error': archive_error})

    # Syntax compilation is a meaningful check of the supplied Java example, not a DB test.
    classes = VALIDATION / 'classes'
    classes.mkdir(exist_ok=True)
    java = subprocess.run(['javac', '-encoding', 'UTF-8', '-d', str(classes), str(ROOT / 'examples' / 'BatchConnectionExample.java')],
                          capture_output=True, text=True, errors='replace', timeout=60)
    check('Java example compilation (JDK 17)', java.returncode == 0,
          {'command': 'javac -encoding UTF-8 -d validation/classes examples/BatchConnectionExample.java',
           'returncode': java.returncode, 'output': java.stdout + java.stderr})
    syntax_errors, text_errors, secret_matches, json_errors = [], [], [], []
    text_files = []
    for folder in ['scripts', 'examples', 'evidence', 'deliverables']:
        text_files.extend(p for p in (ROOT / folder).rglob('*') if p.suffix in {'.py', '.md', '.json', '.sql', '.cli', '.java'})
    for path in text_files:
        content = path.read_text(encoding='utf-8')
        relative = str(path.relative_to(ROOT))
        if '\ufffd' in content or any(re.search(r'[ \t]+$', line) for line in content.splitlines()):
            text_errors.append(relative)
        if re.search(r'\b(?:AKIA|ASIA)[A-Z0-9]{16}\b|-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----', content):
            secret_matches.append(relative)
        if path.suffix == '.py':
            try:
                compile(content, str(path), 'exec')
            except SyntaxError as failure:
                syntax_errors.append(f'{relative}: {failure}')
        if path.suffix == '.json':
            try:
                json.loads(content)
            except json.JSONDecodeError as failure:
                json_errors.append(f'{relative}: {failure}')
    check('Python source syntax', not syntax_errors, syntax_errors)
    check('JSON parsing', not json_errors, json_errors)
    check('UTF-8, replacement characters, trailing whitespace', not text_errors, {'files': len(text_files), 'errors': text_errors})
    check('No common AWS access-key or private-key literals', not secret_matches, secret_matches)
    md_text = MD.read_text(encoding='utf-8')
    missing = []
    for match in re.finditer(r'\]\((\.\./[^)]+)\)', md_text):
        target = (MD.parent / match.group(1)).resolve()
        if target != RESULT.resolve() and not target.exists():
            missing.append(match.group(1))
    check('Markdown local links resolve', not missing, missing)

    report = {
        'generated_at_utc': datetime.now(timezone.utc).isoformat(),
        'result': 'PASS' if all(x['result'] == 'PASS' for x in checks) else 'FAIL',
        'checks': checks,
        'artifacts': [{'file': str(p.relative_to(ROOT)), 'bytes': p.stat().st_size, 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()} for p in [XLSX, MD]],
        'not_executed': ['AWS resource changes', 'DB connection or SQL execution', 'EAP management-model application',
                         'Runtime values on customer DBs', 'Load/latency/failover/restore tests'],
        'git_diff_check': 'Not available: workspace is not a Git repository. Text and artifact checks were performed directly.',
        'visual_review': 'See excel_render_report.json and preview files when present. Visual inspection is reported separately.',
    }
    RESULT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'result': report['result'], 'checks': len(checks), 'failed': [x for x in checks if x['result'] != 'PASS']}, ensure_ascii=True))
    wb.close()
    cached.close()
    return 0 if report['result'] == 'PASS' else 1


if __name__ == '__main__':
    sys.exit(main())
