"""Render representative native cell layouts with Pillow; this is not an Excel renderer."""
import json
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'validation'
XLSX = ROOT / 'deliverables' / 'Aurora_GlobalDatabase_Design_Assessment_2026-10-11.xlsx'
FONT_DIR = Path('C:/Windows/Fonts')
FONT_CACHE = {}


def font(size, bold=False):
    key = (size, bold)
    if key not in FONT_CACHE:
        FONT_CACHE[key] = ImageFont.truetype(str(FONT_DIR / ('meiryob.ttc' if bold else 'meiryo.ttc')),
                                           round(size * 4 / 3), index=2)
    return FONT_CACHE[key]


def rgb(color, default):
    return '#' + color.rgb[-6:] if color and color.type == 'rgb' else default


def wrapped(value, selected_font, available, wrap):
    result = []
    for original in str(value).split('\n'):
        if not wrap:
            result.append(original)
            continue
        line = ''
        for char in original:
            if line and selected_font.getlength(line + char) > available:
                result.append(line)
                line = ''
            line += char
        result.append(line)
    return result


def render(ws, area, filename):
    left, top, right, bottom = range_boundaries(area)
    xs = {left: 12}
    for c in range(left, right + 1):
        dim = ws.column_dimensions.get(get_column_letter(c))
        if not dim:
            dim = next((d for d in ws.column_dimensions.values() if d.min <= c <= d.max), None)
        width = dim.width if dim else 13
        xs[c + 1] = xs[c] + round(width * 7 + 5)
    ys = {top: 12}
    for r in range(top, bottom + 1):
        ys[r + 1] = ys[r] + round((ws.row_dimensions[r].height or 15) * 4 / 3)
    canvas = Image.new('RGB', (xs[right + 1] + 12, ys[bottom + 1] + 12), 'white')
    draw = ImageDraw.Draw(canvas)
    merged = {}
    skip = set()
    for rng in ws.merged_cells.ranges:
        if rng.min_row < top or rng.max_row > bottom or rng.min_col < left or rng.max_col > right:
            continue
        merged[(rng.min_row, rng.min_col)] = (rng.max_row, rng.max_col)
        skip.update((r, c) for r in range(rng.min_row, rng.max_row + 1)
                    for c in range(rng.min_col, rng.max_col + 1) if (r, c) != (rng.min_row, rng.min_col))
    overflows = []
    for r in range(top, bottom + 1):
        for c in range(left, right + 1):
            if (r, c) in skip:
                continue
            end_r, end_c = merged.get((r, c), (r, c))
            cell = ws.cell(r, c)
            x0, x1, y0, y1 = xs[c], xs[end_c + 1], ys[r], ys[end_r + 1]
            fill = rgb(cell.fill.fgColor, '#FFFFFF') if cell.fill.patternType else '#FFFFFF'
            draw.rectangle((x0, y0, x1, y1), fill=fill)
            if cell.border.bottom.style:
                draw.line((x0, y1, x1, y1), fill=rgb(cell.border.bottom.color, '#CCCCCC'), width=1)
            if cell.border.left.style:
                draw.rectangle((x0, y0, x1, y1), outline='#999999', width=1)
            if cell.value is None:
                continue
            selected_font = font(cell.font.sz or 10, cell.font.bold)
            lines = wrapped(cell.value, selected_font, x1-x0-12, cell.alignment.wrap_text)
            ascent, descent = selected_font.getmetrics()
            step = ascent + descent
            height = len(lines) * step
            if height > y1-y0-4 or max(selected_font.getlength(line) for line in lines) > x1-x0-8:
                overflows.append({'cell': cell.coordinate, 'text': str(cell.value)[:90], 'text_height': height, 'cell_height': y1-y0})
            y = y0 + 5
            if cell.alignment.vertical == 'center':
                y = y0 + (y1-y0-height)/2
            color = rgb(cell.font.color, '#242424')
            for line in lines:
                x = x0 + 6
                if cell.alignment.horizontal == 'center':
                    x = x0 + (x1-x0-selected_font.getlength(line))/2
                elif cell.alignment.horizontal == 'right':
                    x = x1-selected_font.getlength(line)-6
                draw.text((x, y), line, font=selected_font, fill=color)
                y += step
    canvas.save(OUT / filename)
    return {'sheet': ws.title, 'range': area, 'png': filename, 'pixels': canvas.size, 'potential_overflows': overflows}


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    samples = [('00_はじめに', 'B2:D26', 'cover_preview.png'),
               ('02_構成図', 'A1:P39', 'architecture_preview.png'),
               ('07_Proxy設定', 'A1:F16', 'proxy_preview.png'),
               ('10_タイムアウト', 'A1:H17', 'timeouts_preview.png'),
               ('20_クラスタ既定', 'A65:M75', 'parameters_preview.png'),
               ('15_接続数試算', 'B2:D29', 'capacity_preview.png')]
    report = {'method': 'Pillow rendering of actual XLSX cell values, merges, widths, heights, fonts, fills and wrapping. Not an Excel screenshot.',
              'samples': [render(wb[s], area, name) for s, area, name in samples]}
    (OUT / 'layout_preview_report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    print(json.dumps(report, ensure_ascii=True))
    wb.close()


if __name__ == '__main__':
    main()
